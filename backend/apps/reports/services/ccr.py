from __future__ import annotations

from datetime import date
from io import BytesIO

from django.db.models import Prefetch
from django.db.models.functions import TruncDate
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from zoneinfo import ZoneInfo

from apps.clinics.ccr_services import get_ccr_settings
from apps.clinics.models import Clinic, ClinicDisabilityType
from apps.encounters.consultation_meta_services import treatment_text_for_encounter
from apps.encounters.models import Diagnosis, Encounter, EncounterConsultationMeta, TreatmentPlan
from apps.patients.models import Patient
from apps.reports.services.common import scope_branch


def _patient_age_label(dob: date, reference: date) -> str:
    months = (reference.year - dob.year) * 12 + (reference.month - dob.month)
    if reference.day < dob.day:
        months -= 1
    if months < 24:
        return f"{months} month" if months == 1 else f"{months} months"
    years = months // 12
    return str(years)


def _gender_label(sex: str) -> str:
    mapping = {
        Patient.Sex.FEMALE: "F",
        Patient.Sex.MALE: "M",
        Patient.Sex.OTHER: "O",
        Patient.Sex.UNKNOWN: "",
    }
    return mapping.get(sex, "")


def _diagnosis_text(encounter: Encounter) -> str:
    rows = Diagnosis.objects.filter(encounter=encounter, status="active").order_by("recorded_at")
    parts = []
    for row in rows:
        text = row.description.strip()
        if row.code:
            text = f"{row.code} {text}".strip()
        parts.append(text)
    return "; ".join(parts)


def _checkbox(value: bool) -> str:
    return "✓" if value else ""


CCR_FORM_ROW_COUNT = 15
CCR_COLUMN_COUNT = 13

_THIN = Side(style="thin")
_ALL_BORDERS = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _format_ccr_date(value: date) -> str:
    return f"{value.day}-{value.month}-{value.year}"


def _mark_checkbox(checked: bool) -> str:
    return "☑" if checked else "☐"


def _population_types_label(header: dict) -> str:
    population_type = header.get("population_type", "")
    return "Types:   " + "   ".join(
        [
            f"{_mark_checkbox(population_type == 'host')} Host",
            f"{_mark_checkbox(population_type == 'idps')} IDPs",
        ]
    )


def _idp_locations_label(header: dict) -> str:
    idp_location = header.get("idp_location", "")
    return "   ".join(
        [
            f"{_mark_checkbox(idp_location == 'forest')} Forest",
            f"{_mark_checkbox(idp_location == 'monastery')} Monastery",
            f"{_mark_checkbox(idp_location == 'school')} School",
            f"{_mark_checkbox(idp_location == 'other')} Other......",
        ]
    )


def _empty_row(no: int) -> dict:
    return {
        "no": no,
        "name": "",
        "age": "",
        "gender": "",
        "diagnosis": "",
        "treatment": "",
        "plw": "",
        "disabled": "",
        "disability_type_number": "",
        "status_old": "",
        "status_new": "",
        "referral": "",
        "referral_cause": "",
        "remark": "",
    }


def _padded_rows(rows: list[dict], minimum: int = CCR_FORM_ROW_COUNT) -> list[dict]:
    padded = list(rows)
    while len(padded) < minimum:
        padded.append(_empty_row(len(padded) + 1))
    return padded


def _disabled_cell(row: dict) -> str:
    disabled = row["disabled"]
    if row["disability_type_number"]:
        return f"{disabled} {row['disability_type_number']}".strip()
    return disabled


def _disability_legend(disability_types: list[dict]) -> str:
    return "Types of disability - " + ", ".join(
        f"{item['sort_order']}. {item['label']}" for item in disability_types
    )


def _header_payload(settings, start: date, end: date) -> dict:
    return {
        "title": "CLINICAL CONSULTATION RECORD [CCR]",
        "project_name": settings.project_name,
        "state_region": settings.state_region,
        "township_mimu": settings.township_mimu,
        "village_mimu": settings.village_mimu,
        "population_type": settings.population_type,
        "idp_location": settings.idp_location,
        "consultant_name": settings.consultant_name,
        "consultant_position": settings.consultant_position,
        "from": start.isoformat(),
        "to": end.isoformat(),
        "report_date": end.isoformat(),
        "report_date_formatted": _format_ccr_date(end),
    }


def _completed_encounters(clinic_id, branch_id, start: date, end: date):
    clinic = Clinic.objects.get(pk=clinic_id)
    zone = ZoneInfo(clinic.timezone)
    queryset = (
        Encounter.objects.filter(clinic_id=clinic_id, status=Encounter.Status.COMPLETED)
        .select_related("patient", "consultation_meta", "consultation_meta__disability_type")
        .prefetch_related(
            Prefetch("diagnosis_entries", queryset=Diagnosis.objects.filter(status="active")),
            Prefetch("treatmentplan_entries", queryset=TreatmentPlan.objects.filter(status="active")),
        )
        .annotate(local_date=TruncDate("completed_at", tzinfo=zone))
        .filter(local_date__gte=start, local_date__lte=end)
        .order_by("completed_at", "created_at")
    )
    return scope_branch(queryset, branch_id)


def _row_payload(index: int, encounter: Encounter, reference: date) -> dict:
    patient = encounter.patient
    meta = getattr(encounter, "consultation_meta", None)
    is_plw = bool(meta and meta.is_plw)
    is_disabled = bool(meta and meta.is_disabled)
    patient_status = meta.patient_status if meta else ""
    referral_made = bool(meta and meta.referral_made)
    disability_number = ""
    if meta and meta.disability_type_id and meta.disability_type:
        disability_number = str(meta.disability_type.sort_order)

    return {
        "no": index,
        "name": patient.full_name,
        "age": _patient_age_label(patient.date_of_birth, reference),
        "gender": _gender_label(patient.sex),
        "diagnosis": _diagnosis_text(encounter),
        "treatment": treatment_text_for_encounter(encounter),
        "plw": _checkbox(is_plw),
        "disabled": _checkbox(is_disabled),
        "disability_type_number": disability_number,
        "status_old": _checkbox(patient_status == EncounterConsultationMeta.PatientStatus.OLD),
        "status_new": _checkbox(patient_status == EncounterConsultationMeta.PatientStatus.NEW),
        "referral": _checkbox(referral_made),
        "referral_cause": (meta.referral_cause if meta else "") or "",
        "remark": (meta.remark if meta else "") or "",
    }


def build_ccr_report(*, clinic_id, branch_id, start: date, end: date) -> dict:
    settings = get_ccr_settings(clinic_id)
    disability_types = list(
        ClinicDisabilityType.objects.filter(clinic_id=clinic_id, is_active=True).order_by("sort_order", "label")
    )
    encounters = list(_completed_encounters(clinic_id, branch_id, start, end))
    reference = end
    rows = [_row_payload(index, encounter, reference) for index, encounter in enumerate(encounters, start=1)]
    return {
        "header": _header_payload(settings, start, end),
        "rows": rows,
        "disability_types": [
            {"sort_order": row.sort_order, "label": row.label} for row in disability_types
        ],
        "row_count": len(rows),
    }


def _style_header_cell(cell, *, bold: bool = False, center: bool = False) -> None:
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=True,
    )
    cell.border = _ALL_BORDERS


def _style_table_cell(cell, *, center: bool = False) -> None:
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="top",
        wrap_text=True,
    )
    cell.border = _ALL_BORDERS


def render_ccr_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CCR"
    header = payload["header"]
    rows = _padded_rows(payload["rows"])
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    column_widths = {
        "A": 4,
        "B": 18,
        "C": 7,
        "D": 7,
        "E": 24,
        "F": 18,
        "G": 5,
        "H": 5,
        "I": 5,
        "J": 5,
        "K": 8,
        "L": 16,
        "M": 14,
    }
    for column, width in column_widths.items():
        ws.column_dimensions[column].width = width

    ws.merge_cells("A1:J1")
    ws["A1"] = header["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws.merge_cells("K1:M1")
    ws["K1"] = "Page: 1"
    ws["K1"].alignment = right
    ws.merge_cells("K2:M2")
    ws["K2"] = f"Date: {header['report_date_formatted']}"
    ws["K2"].alignment = right

    ws.merge_cells("A3:D3")
    ws["A3"] = f"Project Name: {header['project_name']}"
    ws.merge_cells("E3:H3")
    ws["E3"] = f"Township (MIMU): {header['township_mimu']}"
    ws.merge_cells("I3:M3")
    ws["I3"] = _population_types_label(header)

    ws.merge_cells("A4:D4")
    ws["A4"] = f"State/Region: {header['state_region']}"
    ws.merge_cells("E4:H4")
    ws["E4"] = f"Village Name (MIMU): {header['village_mimu']}"
    ws.merge_cells("I4:M4")
    ws["I4"] = _idp_locations_label(header)

    header_row = 5
    subheader_row = 6
    ws.merge_cells(start_row=header_row, start_column=1, end_row=subheader_row, end_column=1)
    ws.merge_cells(start_row=header_row, start_column=2, end_row=subheader_row, end_column=2)
    ws.merge_cells(start_row=header_row, start_column=3, end_row=subheader_row, end_column=3)
    ws.merge_cells(start_row=header_row, start_column=4, end_row=subheader_row, end_column=4)
    ws.merge_cells(start_row=header_row, start_column=5, end_row=subheader_row, end_column=5)
    ws.merge_cells(start_row=header_row, start_column=6, end_row=subheader_row, end_column=6)
    ws.merge_cells(start_row=header_row, start_column=7, end_row=header_row, end_column=8)
    ws.merge_cells(start_row=header_row, start_column=9, end_row=header_row, end_column=10)
    ws.merge_cells(start_row=header_row, start_column=11, end_row=subheader_row, end_column=11)
    ws.merge_cells(start_row=header_row, start_column=12, end_row=subheader_row, end_column=12)
    ws.merge_cells(start_row=header_row, start_column=13, end_row=subheader_row, end_column=13)

    ws.cell(row=header_row, column=1, value="No.")
    ws.cell(row=header_row, column=2, value="Name")
    ws.cell(row=header_row, column=3, value="Age")
    ws.cell(row=header_row, column=4, value="Gender")
    ws.cell(row=header_row, column=5, value="Diagnosis")
    ws.cell(row=header_row, column=6, value="Treatment")
    ws.cell(row=header_row, column=7, value="Vulnerability")
    ws.cell(row=header_row, column=9, value="Status")
    ws.cell(row=header_row, column=11, value="Referral (✓)")
    ws.cell(row=header_row, column=12, value="Cause of Referral")
    ws.cell(row=header_row, column=13, value="Remark")
    ws.cell(row=subheader_row, column=7, value="PLW")
    ws.cell(row=subheader_row, column=8, value="D")
    ws.cell(row=subheader_row, column=9, value="Old")
    ws.cell(row=subheader_row, column=10, value="New")

    for row_index in (header_row, subheader_row):
        for column_index in range(1, CCR_COLUMN_COUNT + 1):
            _style_header_cell(ws.cell(row=row_index, column=column_index), bold=True, center=True)

    data_start_row = subheader_row + 1
    for offset, row in enumerate(rows):
        row_index = data_start_row + offset
        values = [
            row["no"],
            row["name"],
            row["age"],
            row["gender"],
            row["diagnosis"],
            row["treatment"],
            row["plw"],
            _disabled_cell(row),
            row["status_old"],
            row["status_new"],
            row["referral"],
            row["referral_cause"],
            row["remark"],
        ]
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            _style_table_cell(cell, center=column_index in {1, 3, 4, 7, 8, 9, 10, 11})

    footer_row = data_start_row + len(rows) + 1
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
    ws.merge_cells(start_row=footer_row + 1, start_column=1, end_row=footer_row + 1, end_column=8)
    ws.merge_cells(start_row=footer_row + 2, start_column=1, end_row=footer_row + 2, end_column=8)
    ws.cell(row=footer_row, column=1, value="PLW - Pregnant and Lactating Women")
    ws.cell(row=footer_row + 1, column=1, value="D - Disable")
    ws.cell(row=footer_row + 2, column=1, value=_disability_legend(payload["disability_types"]))

    ws.merge_cells(start_row=footer_row, start_column=9, end_row=footer_row, end_column=13)
    ws.merge_cells(start_row=footer_row + 1, start_column=9, end_row=footer_row + 1, end_column=13)
    ws.merge_cells(start_row=footer_row + 2, start_column=9, end_row=footer_row + 2, end_column=13)
    ws.cell(row=footer_row, column=9, value="Signature: _________________________")
    ws.cell(row=footer_row + 1, column=9, value=f"Consultant Name: {header.get('consultant_name', '')}")
    ws.cell(row=footer_row + 2, column=9, value=f"Position: {header.get('consultant_position', '')}")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render_ccr_pdf(payload: dict) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("ccr-normal", parent=styles["Normal"], fontSize=8, leading=10)
    title_style = ParagraphStyle("ccr-title", parent=styles["Heading2"], alignment=1, fontSize=13, spaceAfter=2)
    header = payload["header"]
    rows = _padded_rows(payload["rows"])
    page_width = landscape(A4)[0] - doc.leftMargin - doc.rightMargin

    title_table = Table(
        [
            [Paragraph(header["title"], title_style), ""],
            ["", Paragraph(f"Page: 1<br/>Date: {header['report_date_formatted']}", normal)],
        ],
        colWidths=[page_width * 0.75, page_width * 0.25],
    )
    title_table.setStyle(
        TableStyle(
            [
                ("SPAN", (0, 0), (0, 1)),
                ("ALIGN", (0, 0), (0, 0), "CENTER"),
                ("VALIGN", (0, 0), (0, 0), "MIDDLE"),
                ("ALIGN", (1, 0), (1, 1), "RIGHT"),
                ("VALIGN", (1, 0), (1, 1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )

    meta_table = Table(
        [
            [
                Paragraph(f"Project Name: {header['project_name']}", normal),
                Paragraph(f"Township (MIMU): {header['township_mimu']}", normal),
                Paragraph(_population_types_label(header), normal),
            ],
            [
                Paragraph(f"State/Region: {header['state_region']}", normal),
                Paragraph(f"Village Name (MIMU): {header['village_mimu']}", normal),
                Paragraph(_idp_locations_label(header), normal),
            ],
        ],
        colWidths=[page_width * 0.28, page_width * 0.28, page_width * 0.44],
    )
    meta_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )

    table_data = [
        [
            "No.",
            "Name",
            "Age",
            "Gender",
            "Diagnosis",
            "Treatment",
            "Vulnerability",
            "",
            "Status",
            "",
            "Referral (✓)",
            "Cause of Referral",
            "Remark",
        ],
        ["", "", "", "", "", "", "PLW", "D", "Old", "New", "", "", ""],
    ]
    for row in rows:
        table_data.append(
            [
                str(row["no"]),
                row["name"],
                row["age"],
                row["gender"],
                row["diagnosis"],
                row["treatment"],
                row["plw"],
                _disabled_cell(row),
                row["status_old"],
                row["status_new"],
                row["referral"],
                row["referral_cause"],
                row["remark"],
            ]
        )

    col_widths = [
        page_width * 0.04,
        page_width * 0.12,
        page_width * 0.05,
        page_width * 0.05,
        page_width * 0.16,
        page_width * 0.12,
        page_width * 0.04,
        page_width * 0.04,
        page_width * 0.04,
        page_width * 0.04,
        page_width * 0.05,
        page_width * 0.12,
        page_width * 0.13,
    ]
    data_table = Table(table_data, colWidths=col_widths, repeatRows=2)
    data_table.setStyle(
        TableStyle(
            [
                ("SPAN", (0, 0), (0, 1)),
                ("SPAN", (1, 0), (1, 1)),
                ("SPAN", (2, 0), (2, 1)),
                ("SPAN", (3, 0), (3, 1)),
                ("SPAN", (4, 0), (4, 1)),
                ("SPAN", (5, 0), (5, 1)),
                ("SPAN", (6, 0), (7, 0)),
                ("SPAN", (8, 0), (9, 0)),
                ("SPAN", (10, 0), (10, 1)),
                ("SPAN", (11, 0), (11, 1)),
                ("SPAN", (12, 0), (12, 1)),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
                ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("ALIGN", (0, 0), (-1, 1), "CENTER"),
                ("ALIGN", (0, 2), (0, -1), "CENTER"),
                ("ALIGN", (2, 2), (4, -1), "CENTER"),
                ("ALIGN", (6, 2), (10, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )

    footer_table = Table(
        [
            [
                Paragraph("PLW - Pregnant and Lactating Women", normal),
                Paragraph("Signature: _________________________", normal),
            ],
            [
                Paragraph("D - Disable", normal),
                Paragraph(f"Consultant Name: {header.get('consultant_name', '')}", normal),
            ],
            [
                Paragraph(_disability_legend(payload["disability_types"]), normal),
                Paragraph(f"Position: {header.get('consultant_position', '')}", normal),
            ],
        ],
        colWidths=[page_width * 0.58, page_width * 0.42],
    )
    footer_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )

    story = [
        title_table,
        Spacer(1, 4),
        meta_table,
        Spacer(1, 6),
        data_table,
        Spacer(1, 6),
        footer_table,
    ]
    doc.build(story)
    return buffer.getvalue()
